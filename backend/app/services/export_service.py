from collections.abc import Sequence
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.constants import ORDER_STATUSES, PAYMENT_METHODS
from app.models import Order

MONEY_FORMAT = '"$"#,##0'
DATE_FORMAT = "dd/mm/yyyy hh:mm"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4F46E5")

ORDER_COLUMNS = [
    ("Pedido", 14), ("Fecha", 17), ("Estado", 12), ("Cliente", 28), ("Celular", 15),
    ("Dirección", 36), ("Ciudad", 16), ("Método de pago", 16), ("Productos", 50),
    ("Subtotal", 13), ("Envío", 12), ("Total", 13), ("Observaciones", 40),
]
ITEM_COLUMNS = [
    ("Pedido", 14), ("Fecha", 17), ("Estado", 12), ("Cliente", 28), ("Producto", 36),
    ("Cantidad", 10), ("Precio unitario", 15), ("Subtotal", 13),
]
ORDER_MONEY_COLUMNS = {10, 11, 12}
ITEM_MONEY_COLUMNS = {7, 8}


def _local(dt: datetime, tz: ZoneInfo) -> datetime:
    # Excel no admite zonas horarias: se guarda la hora local del negocio.
    return dt.astimezone(tz).replace(tzinfo=None)


def _write_row(ws: Worksheet, values: list, money_columns: set[int]) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        # openpyxl interpreta los textos que empiezan por "=" como fórmulas.
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"
        if cell.column in money_columns:
            cell.number_format = MONEY_FORMAT
        elif isinstance(cell.value, datetime):
            cell.number_format = DATE_FORMAT


def _setup_sheet(ws: Worksheet, columns: list[tuple[str, int]]) -> None:
    ws.append([name for name, _ in columns])
    for index, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
        header = ws.cell(row=1, column=index)
        header.font = HEADER_FONT
        header.fill = HEADER_FILL
    ws.freeze_panes = "A2"


def build_orders_workbook(orders: Sequence[Order], tz: ZoneInfo) -> bytes:
    """Excel con una hoja de pedidos (una fila por pedido) y otra con el detalle de productos."""
    wb = Workbook()
    orders_ws = wb.active
    orders_ws.title = "Pedidos"
    items_ws = wb.create_sheet("Productos")
    _setup_sheet(orders_ws, ORDER_COLUMNS)
    _setup_sheet(items_ws, ITEM_COLUMNS)

    for order in orders:
        created_at = _local(order.created_at, tz)
        status = ORDER_STATUSES.get(order.status, order.status)
        products = ", ".join(f"{i.quantity} × {i.product_name}" for i in order.items)
        _write_row(
            orders_ws,
            [
                order.order_number, created_at, status, order.customer.name, order.customer.phone,
                order.shipping_address, order.shipping_city,
                PAYMENT_METHODS.get(order.payment_method, order.payment_method), products,
                order.subtotal, order.shipping_cost, order.total, order.notes,
            ],
            ORDER_MONEY_COLUMNS,
        )
        for item in order.items:
            _write_row(
                items_ws,
                [
                    order.order_number, created_at, status, order.customer.name, item.product_name,
                    item.quantity, item.unit_price, item.subtotal,
                ],
                ITEM_MONEY_COLUMNS,
            )

    for ws in (orders_ws, items_ws):
        ws.auto_filter.ref = ws.dimensions

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
