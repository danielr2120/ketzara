from concurrent.futures import ThreadPoolExecutor
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from tests.conftest import order_payload


def test_create_order_calculates_totals_and_number(client, product):
    res = client.post("/api/orders", json=order_payload(product["id"], shipping_cost=8000))
    assert res.status_code == 201, res.text
    order = res.json()
    assert order["order_number"] == "PED-000001"
    assert order["status"] == "pending"
    assert order["subtotal"] == 50000
    assert order["shipping_cost"] == 8000
    assert order["total"] == 58000
    assert order["customer"]["phone"] == "3001234567"
    assert order["shipping_address"] == "Calle 1 # 2-3"
    item = order["items"][0]
    assert item["product_name"] == "Producto A"
    assert item["unit_price"] == 25000

    second = client.post("/api/orders", json=order_payload(product["id"])).json()
    assert second["order_number"] == "PED-000002"
    assert second["customer"]["id"] == order["customer"]["id"]


def test_custom_price_does_not_change_product(client, product):
    payload = order_payload(items=[{"product_id": product["id"], "quantity": 3, "unit_price": 20000}])
    order = client.post("/api/orders", json=payload).json()
    assert order["total"] == 60000
    products = client.get("/api/products").json()
    assert products[0]["price"] == 25000


def test_free_text_item(client):
    payload = order_payload(items=[{"product_name": "Servicio especial", "quantity": 1, "unit_price": 15000}])
    res = client.post("/api/orders", json=payload)
    assert res.status_code == 201
    assert res.json()["items"][0]["product_id"] is None


def test_history_kept_after_product_change(client, product):
    order = client.post("/api/orders", json=order_payload(product["id"])).json()
    client.put(f"/api/products/{product['id']}", json={"name": "Renombrado", "price": 99999})
    stored = client.get(f"/api/orders/{order['id']}").json()
    assert stored["items"][0]["product_name"] == "Producto A"
    assert stored["items"][0]["unit_price"] == 25000


def test_order_validation(client, product):
    payload = order_payload(
        items=[{"product_id": product["id"], "quantity": 0}],
        payment_method="bitcoin",
        shipping_cost=-5,
    )
    payload["customer"]["name"] = ""
    res = client.post("/api/orders", json=payload)
    assert res.status_code == 422
    fields = {e["field"] for e in res.json()["errors"]}
    assert {"customer.name", "items.0.quantity", "payment_method", "shipping_cost"} <= fields


def test_inactive_or_missing_product_rejected(client, product):
    client.put(f"/api/products/{product['id']}", json={"name": "Producto A", "price": 1, "active": False})
    res = client.post("/api/orders", json=order_payload(product["id"]))
    assert res.status_code == 422
    res = client.post("/api/orders", json=order_payload(9999))
    assert res.status_code == 422


def test_update_order_recalculates(client, product):
    order = client.post("/api/orders", json=order_payload(product["id"])).json()
    payload = order_payload(
        items=[
            {"product_id": product["id"], "product_name": "Producto A", "quantity": 1, "unit_price": 25000},
            {"product_name": "Extra", "quantity": 2, "unit_price": 1000},
        ],
        shipping_cost=5000,
        payment_method="nequi",
        notes="Entregar en la tarde",
    )
    res = client.put(f"/api/orders/{order['id']}", json=payload)
    assert res.status_code == 200, res.text
    updated = res.json()
    assert updated["order_number"] == order["order_number"]
    assert updated["subtotal"] == 27000
    assert updated["total"] == 32000
    assert len(updated["items"]) == 2
    assert updated["payment_method"] == "nequi"


def test_status_change_and_cancelled_cannot_be_edited(client, product):
    order = client.post("/api/orders", json=order_payload(product["id"])).json()
    res = client.patch(f"/api/orders/{order['id']}/status", json={"status": "shipped"})
    assert res.json()["status"] == "shipped"

    bad = client.patch(f"/api/orders/{order['id']}/status", json={"status": "perdido"})
    assert bad.status_code == 422

    client.patch(f"/api/orders/{order['id']}/status", json={"status": "cancelled"})
    res = client.put(f"/api/orders/{order['id']}", json=order_payload(product["id"]))
    assert res.status_code == 422
    assert res.json()["detail"] == "No se puede editar un pedido cancelado"


def test_search_and_filters(client, product):
    client.post("/api/orders", json=order_payload(product["id"]))
    other = order_payload(product["id"])
    other["customer"] = {"name": "María López", "phone": "3109998877", "address": "Cra 7"}
    second = client.post("/api/orders", json=other).json()
    client.patch(f"/api/orders/{second['id']}/status", json={"status": "cancelled"})

    def search(**params):
        return client.get("/api/orders", params=params).json()

    assert search()["total"] == 2
    assert search()["sales_total"] == 50000
    assert [o["customer_name"] for o in search(search="maría")["items"]] == ["María López"]
    assert search(search="999 88")["items"][0]["order_number"] == "PED-000002"
    assert search(search="PED-2")["items"][0]["order_number"] == "PED-000002"
    assert search(search="000001")["items"][0]["order_number"] == "PED-000001"
    assert search(status="cancelled")["total"] == 1
    assert search(search="%")["total"] == 0

    today = date.today().isoformat()
    assert search(date_from=today, date_to=today)["total"] == 2
    assert search(date_from="2000-01-01", date_to="2000-01-02")["total"] == 0
    assert client.get("/api/orders", params={"status": "nope"}).status_code == 422


def test_order_not_found(client):
    assert client.get("/api/orders/999").status_code == 404


def test_lookup_by_number(client, product):
    order = client.post("/api/orders", json=order_payload(product["id"])).json()
    for number in ["PED-000001", "ped-1", " 1 ", "000001"]:
        res = client.get(f"/api/orders/number/{number}")
        assert res.status_code == 200, number
        assert res.json()["id"] == order["id"]
    missing = client.get("/api/orders/number/PED-000099")
    assert missing.status_code == 404
    assert missing.json()["detail"] == "No encontramos un pedido con ese número"
    assert client.get("/api/orders/number/abc").status_code == 404


def test_public_tracking_link(client, product):
    first = client.post("/api/orders", json=order_payload(product["id"], shipping_cost=5000)).json()
    second = client.post("/api/orders", json=order_payload(product["id"])).json()
    assert len(first["tracking_code"]) >= 12
    assert first["tracking_code"] != second["tracking_code"]

    res = client.get(f"/api/public/orders/{first['tracking_code']}")
    assert res.status_code == 200
    data = res.json()
    assert data["order_number"] == "PED-000001"
    assert data["status_label"] == "Pendiente"
    assert data["total"] == 55000
    assert "cancelled" not in [s["value"] for s in data["steps"]]
    assert set(data) == {
        "order_number", "created_at", "updated_at", "status", "status_label",
        "steps", "items", "shipping_cost", "total",
    }
    assert set(data["items"][0]) == {"product_name", "quantity", "subtotal"}

    assert client.get("/api/public/orders/no-existe").status_code == 404


def test_public_catalog_hides_inactive_and_sold_out(client, product):
    client.post("/api/products", json={"name": "Inactivo", "price": 1000, "active": False})
    client.post("/api/products", json={"name": "Agotado", "price": 1000, "stock": 0})
    client.post("/api/products", json={"name": "Sin control de stock", "price": 1000})

    products = client.get("/api/public/products").json()
    assert [p["name"] for p in products] == ["Producto A", "Sin control de stock"]
    assert set(products[0]) == {"id", "name", "description", "price"}


def test_public_order_uses_catalog_price(client, product):
    payload = {
        "customer": {"name": "Ana Ruiz", "phone": "311 555 0000", "address": "Calle 9 # 8-7"},
        "items": [{"product_id": product["id"], "quantity": 3}],
        "payment_method": "nequi",
        "notes": "Timbre dañado",
    }
    res = client.post("/api/public/orders", json=payload)
    assert res.status_code == 201, res.text
    created = res.json()
    assert created["order_number"] == "PED-000001"
    assert created["status"] == "pending"
    assert created["total"] == 75000
    assert created["shipping_cost"] == 0
    assert "customer" not in created

    tracked = client.get(f"/api/public/orders/{created['tracking_code']}").json()
    assert tracked["order_number"] == "PED-000001"

    order = client.get("/api/orders/number/PED-000001").json()
    assert order["customer"]["phone"] == "3115550000"
    assert order["notes"] == "Timbre dañado"
    assert order["items"][0]["unit_price"] == 25000


def test_public_order_rejects_price_override_and_unavailable_products(client, product):
    base = {
        "customer": {"name": "Ana Ruiz", "phone": "3115550000", "address": "Calle 9"},
        "payment_method": "cash",
    }
    res = client.post(
        "/api/public/orders",
        json={**base, "items": [{"product_id": product["id"], "quantity": 1, "unit_price": 1}]},
    )
    assert res.status_code == 422
    res = client.post("/api/public/orders", json={**base, "shipping_cost": 0, "items": [{"product_id": product["id"], "quantity": 1}]})
    assert res.status_code == 422

    sold_out = client.post("/api/products", json={"name": "Agotado", "price": 1000, "stock": 0}).json()
    res = client.post("/api/public/orders", json={**base, "items": [{"product_id": sold_out["id"], "quantity": 1}]})
    assert res.status_code == 422
    assert res.json()["detail"] == "El producto «Agotado» está agotado"

    client.put(f"/api/products/{product['id']}", json={"name": "Producto A", "price": 1, "active": False})
    res = client.post("/api/public/orders", json={**base, "items": [{"product_id": product["id"], "quantity": 1}]})
    assert res.status_code == 422


def test_export_orders_to_excel(client, product):
    client.post("/api/orders", json=order_payload(product["id"], shipping_cost=5000, notes="Urgente"))
    other = order_payload(product["id"])
    other["customer"] = {"name": "=HYPERLINK(\"http://x\")", "phone": "3109998877", "address": "Cra 7"}
    second = client.post("/api/orders", json=other).json()
    client.patch(f"/api/orders/{second['id']}/status", json={"status": "cancelled"})

    res = client.get("/api/orders/export")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in res.headers["content-disposition"]

    wb = load_workbook(BytesIO(res.content))
    orders = list(wb["Pedidos"].iter_rows(values_only=True))
    assert orders[0][:3] == ("Pedido", "Fecha", "Estado")
    assert [row[0] for row in orders[1:]] == ["PED-000002", "PED-000001"]
    first = orders[2]
    assert first[2] == "Pendiente"
    assert first[4] == "3001234567"
    assert first[7] == "Efectivo"
    assert first[8] == "2 × Producto A"
    assert first[9:13] == (50000, 5000, 55000, "Urgente")
    assert wb["Pedidos"]["D2"].data_type == "s"
    assert orders[1][3] == '=HYPERLINK("http://x")'

    items = list(wb["Productos"].iter_rows(values_only=True))
    assert len(items) == 3
    assert items[1][4:] == ("Producto A", 2, 25000, 50000)

    filtered = load_workbook(BytesIO(client.get("/api/orders/export", params={"status": "cancelled"}).content))
    assert [row[0] for row in filtered["Pedidos"].iter_rows(min_row=2, values_only=True)] == ["PED-000002"]


def test_health(client):
    assert client.get("/api/health").json() == {"status": "ok"}


def test_concurrent_orders_get_unique_numbers(client, product):
    def create(i: int):
        payload = order_payload(product["id"])
        payload["customer"]["phone"] = f"30000000{i:02d}"
        return client.post("/api/orders", json=payload)

    with ThreadPoolExecutor(max_workers=8) as pool:
        responses = list(pool.map(create, range(16)))
    assert all(r.status_code == 201 for r in responses)
    numbers = {r.json()["order_number"] for r in responses}
    assert len(numbers) == 16


def test_dashboard_summary(client, product):
    client.post("/api/orders", json=order_payload(product["id"], shipping_cost=10000))
    second = client.post("/api/orders", json=order_payload(product["id"])).json()
    client.patch(f"/api/orders/{second['id']}/status", json={"status": "delivered"})
    third = client.post("/api/orders", json=order_payload(product["id"])).json()
    client.patch(f"/api/orders/{third['id']}/status", json={"status": "cancelled"})

    summary = client.get("/api/dashboard/summary").json()
    assert summary["orders_today"] == 3
    assert summary["pending_orders"] == 1
    assert summary["delivered_this_month"] == 1
    assert summary["sales_today"] == 110000
    assert summary["sales_this_month"] == 110000
    assert len(summary["recent_orders"]) == 3
    assert summary["recent_orders"][0]["order_number"] == "PED-000003"


def test_meta(client):
    meta = client.get("/api/meta").json()
    assert {"value": "nequi", "label": "Nequi"} in meta["payment_methods"]
    assert [s["value"] for s in meta["order_statuses"]][0] == "pending"
